"""Shows the row of the document a number comes from, highlighted."""
from pathlib import Path

import pymupdf
import streamlit as st
from openpyxl import load_workbook

HIGHLIGHT = (1.0, 0.72, 0.30)  # the ochre of the theme
ROWS_AROUND = 2  # rows of context shown above and below the cited row


def show_source(path: Path, source, caption: str):
    st.caption(f"{caption} · {source.document}" + (f" · page {source.page}" if source.page else "") + f" · row {source.row}")
    if path.suffix.lower() == ".pdf":
        st.image(pdf_row_image(path, source.page or 1, source.raw_text), width="stretch")
    else:
        st.dataframe(sheet_rows(path, int(source.row)), hide_index=True, width="stretch")


def pdf_row_image(path: Path, page_number: int, raw_text: str) -> bytes:
    """A picture of the part of the page around the cited row, with the row highlighted."""
    document = pymupdf.open(path)
    page = document[page_number - 1]
    anchor = raw_text.split()[0]  # the supplier code, or the first word of a footnote
    found = page.search_for(anchor)
    if not found:
        return page.get_pixmap(dpi=110).tobytes("png")
    row = found[0]
    band = pymupdf.Rect(page.rect.x0 + 20, row.y0 - 2, page.rect.x1 - 20, row.y1 + 2)
    page.draw_rect(band, color=HIGHLIGHT, fill=HIGHLIGHT, fill_opacity=0.35, width=0.8)
    clip = pymupdf.Rect(page.rect.x0, max(page.rect.y0, row.y0 - 70), page.rect.x1, min(page.rect.y1, row.y1 + 70))
    return page.get_pixmap(dpi=150, clip=clip).tobytes("png")


def sheet_rows(path: Path, row_number: int):
    """The cited spreadsheet row with a few rows around it. The cited one is marked with an arrow."""
    sheet = load_workbook(path, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    header = next(row for row in rows if row and str(row[0] or "").strip() == "Product Code")
    table = []
    for number in range(max(1, row_number - ROWS_AROUND), min(len(rows), row_number + ROWS_AROUND) + 1):
        values = rows[number - 1]
        if values == header or not values[0]:
            continue
        line = {"": "➜" if number == row_number else "", "row": number}
        line.update({str(name): value for name, value in zip(header, values) if name})
        table.append(line)
    return table
