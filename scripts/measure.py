#!/usr/bin/env python3
"""Measures, with a date: how Margin Guardian does against a plain chatbot, and how long the AI calls take.

Three parts:

1. BASELINE. A generic LLM (Gemma itself, without our pipeline) gets the old and the new
   price list of each demo supplier as plain text and is asked which prices went up and by
   how much. Its answer and the pipeline's answer (mg/report.py, the chain the app uses) are
   both scored against data/demo/ANSWER_KEY.md: increases found, missed, invented, and found
   with the right size (within 0.5 percentage points of the true per-unit change).

2. TIMINGS of the AI calls the product makes: the footnote reading on Gemma (Token Factory),
   the PDF reading through the Evolus extractor, the email text from the Evolus email tool.

3. THE PDF GUARD. Every live PDF reading is checked against the PDF text by plain code
   (mg/extract/pdf.py, check_against_pdf). We count how many readings the guard stopped.

Nothing is saved over the demo readings. The results go to data/measurements/<date and time>.json.

Usage:  .venv/bin/python scripts/measure.py            (all parts, about 5-10 minutes)
        .venv/bin/python scripts/measure.py --no-evolus (only Token Factory calls)
"""
import argparse
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from mg import llm  # noqa: E402
from mg.extract import pdf  # noqa: E402
from mg.negotiate import brief_for, call_email_tool, template_problem  # noqa: E402
from mg.report import DEMO_SUPPLIERS, SUPPLIERS_DIR, run_supplier  # noqa: E402

OUT_DIR = ROOT / "data" / "measurements"
TOLERANCE_POINTS = 0.5  # a size counts as right if within 0.5 percentage points
PAUSE_SECONDS = 4  # between calls: Token Factory allows 20 requests a minute

# The answer key (data/demo/ANSWER_KEY.md): supplier code -> true change of the UNIT cost, in %.
# Only what makes a product cost more. The decrease, the new product, the dropped product and
# the new minimum order are not increases: a tool that lists them as increases has invented them.
EXPECTED = {
    "veloparts": {"VP-10233": 5.98, "VP-10234": 3.07, "VP-30108": 3.72, "VP-60030": 4.22},
    "ciclo": {
        "CC-S305": 7.99, "CC-C705": 4.57,
        # the footnote: +4% on the aluminium items, marked with *
        "CC-A118": 4.0, "CC-A121": 4.0, "CC-A130": 4.0, "CC-A140": 4.0, "CC-M520": 4.0, "CC-B610": 4.0,
    },
    "trailgear": {
        "TG-T1001": 5.0, "TG-T1002": 5.0, "TG-T1010": 5.02, "TG-T1020": 5.04, "TG-L5001": 4.26,
        "TG-B4001": 6.25,  # the pack trap: pack of 4 at 9.60 -> pack of 2 at 5.10
    },
}
HIDDEN = {"CC-A118", "CC-A121", "CC-A130", "CC-A140", "CC-M520", "CC-B610", "TG-B4001"}

BASELINE_SYSTEM = "You are a helpful assistant for a small online shop. Answer with JSON only."
BASELINE_QUESTION = (
    "Below are the OLD and the NEW price list of the same supplier. "
    "Which products cost us more in the new price list, and by how much in percent? "
    "Include every product that costs more for any reason. "
    'Answer {"increases": [{"code": "supplier product code", "percent": number}]}.\n\n'
    "OLD PRICE LIST:\n{old}\n\nNEW PRICE LIST:\n{new}"
)


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def price_list_text(path: Path) -> str:
    """The price list as plain text, the way someone would paste it into a chatbot."""
    if path.suffix == ".pdf":
        return "\n".join(pdf.compact(text) for _, _, text in pdf.pdf_text_rows(path))
    rows = load_workbook(path, data_only=True).active.iter_rows(values_only=True)
    return "\n".join(" | ".join("" if cell is None else str(cell) for cell in row) for row in rows)


def score(claimed: dict, expected: dict) -> dict:
    """claimed and expected: supplier code -> percent. Returns the counts and the codes."""
    found = [code for code in expected if code in claimed]
    right_size = [code for code in found if abs(float(claimed[code]) - expected[code]) <= TOLERANCE_POINTS]
    return {
        "expected": len(expected),
        "found": len(found),
        "found_right_size": len(right_size),
        "missed": sorted(code for code in expected if code not in claimed),
        "invented": sorted(code for code in claimed if code not in expected),
        "hidden_expected": len(HIDDEN & set(expected)),
        "hidden_found": len(HIDDEN & set(found)),
    }


def total(scores: list[dict]) -> dict:
    """The scores of the three suppliers added up."""
    keys = ("expected", "found", "found_right_size", "hidden_expected", "hidden_found")
    result = {key: sum(s[key] for s in scores) for key in keys}
    result["missed"] = sum(len(s["missed"]) for s in scores)
    result["invented"] = sum(len(s["invented"]) for s in scores)
    return result


def pipeline_claims(key: str) -> dict:
    """What Margin Guardian says costs more, on the saved demo readings: code -> percent."""
    report = run_supplier(key)
    return {
        report.products[finding.sku].supplier_code: finding.pct_change
        for finding in report.findings
        if finding.kind in ("price_up", "pack_change", "surcharge_note")
    }


def baseline_claims(key: str) -> tuple[dict, float]:
    """What a generic LLM says costs more: code -> percent, and how long it took."""
    _, old_name, new_name = DEMO_SUPPLIERS[key]
    question = BASELINE_QUESTION.replace("{old}", price_list_text(SUPPLIERS_DIR / "old" / old_name))
    question = question.replace("{new}", price_list_text(SUPPLIERS_DIR / "new" / new_name))
    start = time.monotonic()
    answer = llm.ask_json(question, system=BASELINE_SYSTEM, max_tokens=2000)
    seconds = time.monotonic() - start
    claims = {}
    for item in answer.get("increases") or []:
        try:
            claims[str(item["code"]).strip()] = float(item["percent"])
        except (KeyError, TypeError, ValueError):
            claims[str(item)] = 0.0  # an entry without a usable code or percent still counts as claimed
    return claims, seconds


def part_baseline(runs: int) -> dict:
    print("1. Pipeline vs a generic LLM")
    pipeline = {key: score(pipeline_claims(key), EXPECTED[key]) for key in DEMO_SUPPLIERS}
    print("   pipeline:", total(list(pipeline.values())))

    baseline_runs = []
    for run in range(1, runs + 1):
        per_supplier, seconds = {}, {}
        for key in DEMO_SUPPLIERS:
            try:
                claims, seconds[key] = baseline_claims(key)
                per_supplier[key] = score(claims, EXPECTED[key])
                per_supplier[key]["claimed"] = claims
            except llm.LLMError as error:
                per_supplier[key] = {"error": str(error)[:200]}
            time.sleep(PAUSE_SECONDS)
        ok = [s for s in per_supplier.values() if "error" not in s]
        summary = total(ok) if len(ok) == len(per_supplier) else None
        baseline_runs.append({"run": run, "at": now(), "total": summary, "seconds": seconds, "suppliers": per_supplier})
        print(f"   generic LLM, run {run}:", summary or per_supplier)
    return {
        "pipeline": {"note": "on the saved demo readings; the live reading is measured in part 3",
                     "total": total(list(pipeline.values())), "suppliers": pipeline},
        "generic_llm": {"model": llm.DEFAULT_MODEL, "input": "old and new price list as plain text",
                        "question": BASELINE_QUESTION.split("\n\n")[0], "runs": baseline_runs},
    }


def part_footnote(runs: int) -> dict:
    print("2. Footnote reading on Token Factory")
    note = json.loads((pdf.config.EXTRACT_RESULTS_DIR / "ciclo-componenti_listino_2026-10.pdf.json").read_text())
    text = note["data"]["notes"][1]["text"]  # the aluminium surcharge
    calls = []
    for _ in range(runs):
        start = time.monotonic()
        try:
            answer = pdf.read_note(text)
            calls.append({"seconds": round(time.monotonic() - start, 1), "answer": answer})
        except llm.LLMError as error:
            calls.append({"seconds": round(time.monotonic() - start, 1), "error": str(error)[:200]})
        print("  ", calls[-1])
        time.sleep(PAUSE_SECONDS)
    return {"model": llm.DEFAULT_MODEL, "expected": {"surcharge_pct": 4.0, "valid_from": "2026-11-01"}, "calls": calls}


def part_pdf_guard() -> dict:
    print("3. PDF reading through Evolus, checked by the guard")
    readings = []
    for folder, name in [("old", v[1]) for v in DEMO_SUPPLIERS.values()] + [("new", v[2]) for v in DEMO_SUPPLIERS.values()]:
        if not name.endswith(".pdf"):
            continue
        path = SUPPLIERS_DIR / folder / name
        rows = pdf.pdf_text_rows(path)
        priced_rows = sum(1 for _, _, text in rows if pdf.PRICE_PATTERN.search(text))
        start = time.monotonic()
        try:
            data = pdf.call_evolus(path)
            problems = pdf.check_against_pdf(data, rows)
            reading = {"lines_read": len(data["lines"]), "problems": problems}
        except pdf.ExtractionError as error:
            reading = {"error": str(error)[:200]}
        reading = {"document": name, "priced_rows_in_pdf": priced_rows,
                   "seconds": round(time.monotonic() - start, 1), **reading,
                   "stopped_by_guard": bool(reading.get("problems"))}
        readings.append(reading)
        print("  ", reading)
        time.sleep(PAUSE_SECONDS)
    return {"extractor_model": pdf.extractor_model(), "readings": readings,
            "stopped_by_guard": sum(r["stopped_by_guard"] for r in readings),
            "failed_calls": sum("error" in r for r in readings)}


def part_email() -> dict:
    print("4. Email text from the Evolus email tool")
    calls = []
    for supplier, _, _ in DEMO_SUPPLIERS.values():
        start = time.monotonic()
        try:
            text = call_email_tool(brief_for(supplier))
            calls.append({"supplier": supplier, "seconds": round(time.monotonic() - start, 1),
                          "usable": template_problem(text) is None, "problem": template_problem(text)})
        except Exception as error:  # noqa: BLE001  any failure is a result here
            calls.append({"supplier": supplier, "seconds": round(time.monotonic() - start, 1), "error": str(error)[:200]})
        print("  ", calls[-1])
        time.sleep(PAUSE_SECONDS)
    return {"calls": calls}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--runs", type=int, default=3, help="how many times the generic LLM is asked (default 3)")
    parser.add_argument("--no-evolus", action="store_true", help="skip the Evolus calls (PDF reading and email)")
    args = parser.parse_args()

    result = {"started_at": now()}
    result["baseline"] = part_baseline(args.runs)
    result["footnote"] = part_footnote(3)
    if not args.no_evolus:
        result["pdf_guard"] = part_pdf_guard()
        result["email"] = part_email()
    result["finished_at"] = now()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = result["started_at"][:16].replace(":", "")
    json_path = OUT_DIR / f"{stamp}.json"
    json_path.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")
    print("written:", json_path.relative_to(ROOT))


if __name__ == "__main__":
    main()
