"""Reader for Excel price lists. Plain code, no model involved.

Expected layout (the TrailGear files in data/demo/suppliers):
  row 1      title, starting with the supplier name: "TrailGear Ltd ... - Trade price list ..."
  next rows  free-text notes, e.g. "All prices in GBP, ex VAT."
  one row    the column headers: Product Code | Description | ... | Pack Qty | Trade Price (GBP) | ...
  then       one product per row
"""
import re
from pathlib import Path

from openpyxl import load_workbook

from mg.schemas import DocumentNote, PriceLine, SourceRef

CODE_HEADER = "Product Code"
DESCRIPTION_HEADER = "Description"
PACK_HEADER = "Pack Qty"
PRICE_HEADER = "Trade Price"  # the real header also carries the currency: "Trade Price (GBP)"


def read_xlsx_price_list(path: Path) -> tuple[list[PriceLine], list[DocumentNote]]:
    sheet = load_workbook(path, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))  # rows[0] is Excel row 1

    header_index = find_header_row(rows)
    headers = [str(cell or "").strip() for cell in rows[header_index]]
    code_col = headers.index(CODE_HEADER)
    description_col = headers.index(DESCRIPTION_HEADER)
    pack_col = headers.index(PACK_HEADER)
    price_col = next(i for i, header in enumerate(headers) if header.startswith(PRICE_HEADER))

    title = str(rows[0][0])
    supplier = title.split(" - ")[0].replace("(fictional)", "").strip()
    note_rows = [(index + 1, row_text(rows[index])) for index in range(1, header_index)]
    note_rows = [(number, text) for number, text in note_rows if text]
    currency = find_currency(headers[price_col], [text for _, text in note_rows])

    def source(row_number, text, pack_size=None):
        price_row = {"currency": currency, "pack_size": pack_size} if pack_size else {}
        return SourceRef(document=path.name, sheet=sheet.title, row=row_number, raw_text=text, **price_row)

    notes = [DocumentNote(text=text, source=source(number, text)) for number, text in note_rows]

    lines = []
    for index in range(header_index + 1, len(rows)):
        row = rows[index]
        if not row[code_col]:
            continue  # empty row
        lines.append(
            PriceLine(
                supplier=supplier,
                supplier_code=str(row[code_col]).strip(),
                description=str(row[description_col]).strip(),
                price=float(row[price_col]),
                currency=currency,
                pack_size=int(row[pack_col]),
                source=source(index + 1, row_text(row), int(row[pack_col])),
            )
        )
    return lines, notes


def find_header_row(rows) -> int:
    """Index of the row that holds the column headers."""
    for index, row in enumerate(rows):
        if row and str(row[0] or "").strip() == CODE_HEADER:
            return index
    raise ValueError(f"no header row starting with '{CODE_HEADER}'")


def find_currency(price_header: str, notes: list[str]) -> str:
    """The currency as the file declares it: in the price header, or else in a note."""
    in_header = re.search(r"\(([A-Z]{3})\)", price_header)
    if in_header:
        return in_header.group(1)
    for note in notes:
        in_note = re.search(r"prices in ([A-Z]{3})\b", note)
        if in_note:
            return in_note.group(1)
    raise ValueError("the price list does not say which currency it uses")


def row_text(row) -> str:
    """The row as written in the sheet, cells joined by ' | '."""
    return " | ".join(str(cell) for cell in row if cell is not None)
